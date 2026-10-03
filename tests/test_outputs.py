from openpyxl import load_workbook

from cursus.cli import main
from cursus.datavis import COLUMNS, rows, write_xlsx
from cursus.mermaid import to_mermaid
from cursus.prompts import reading_prompt


def test_excel_table_has_visios_columns(expense, tmp_path):
    ws = load_workbook(write_xlsx(expense, tmp_path / "flow.xlsx")).active
    assert [c.value for c in ws[1]] == COLUMNS
    assert "ProcessData" in ws.tables
    table = {r[0]: r for r in rows(expense)}
    assert table["d1"][2] == "s4,s5" and table["d1"][3] == "Yes,No"
    assert table["d1"][4] == "Decision" and table["start"][4] == "Start"
    assert table["s3"][5] == "Line manager"
    assert table["d3"][2].endswith("d3_not_stated") and table["d3_not_stated"][7] == "Not stated in the text"


def test_commas_in_labels_do_not_add_phantom_arrows(expense):
    expense.flows[4].label = "Yes, fully"
    assert {r[0]: r for r in rows(expense)}["d1"][3] == "Yes; fully,No"


def test_mermaid_has_lanes_labels_and_the_dashed_stub(expense):
    text = to_mermaid(expense)
    assert text.startswith("flowchart LR")
    assert 'subgraph n_manager["Line manager"]' in text
    assert 'n_d1 -- "Yes" --> n_s4' in text
    assert "class n_d3_not_stated assumed" in text


def test_prompt_carries_the_text_and_the_layout(expense_source):
    prompt = reading_prompt(expense_source)
    assert "An employee fills in the claim form" in prompt
    assert '"from"' in prompt and '"quote"' in prompt
    assert "do not invent the missing path" in prompt


def test_build_writes_four_files(tmp_path, capsys):
    code = main(["build", "examples/expense-claim/reading.json", "--source", "examples/expense-claim/process.txt", "--out", str(tmp_path), "--name", "x"])
    assert code == 0
    assert {p.name for p in tmp_path.iterdir()} == {"x.vsdx", "x.xlsx", "x.svg", "x.md"}
    md = (tmp_path / "x.md").read_text()
    assert "What happens when the finance director does not sign off a claim?" in md
    assert "| Check claim against policy | Line manager |" in md


def test_build_refuses_a_reading_whose_quotes_are_not_in_the_text(tmp_path, capsys):
    other = tmp_path / "other.txt"
    other.write_text("A completely different text.")
    code = main(["build", "examples/expense-claim/reading.json", "--source", str(other), "--out", str(tmp_path / "out")])
    assert code == 1
    assert not (tmp_path / "out").exists()
    assert "quote-not-found" in capsys.readouterr().out


def test_reading_wrapped_in_a_code_fence_still_loads(tmp_path):
    raw = open("examples/hello/reading.json").read()
    wrapped = tmp_path / "r.json"
    wrapped.write_text("Here you go:\n```json\n" + raw + "\n```\n")
    assert main(["check", str(wrapped)]) == 0
    empty = tmp_path / "e.json"
    empty.write_text("no json here")
    assert main(["check", str(empty)]) == 2
